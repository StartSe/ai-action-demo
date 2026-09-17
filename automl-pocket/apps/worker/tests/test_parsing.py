"""Testes de leitura dos 3 formatos suportados (csv, xlsx, json)."""

import json
from pathlib import Path

import pandas as pd
import pytest

from jobs.parsing import ParseError, read_dataset_file


def test_csv_virgula_utf8(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("nome,idade\nAna,34\nJoão,28\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["nome", "idade"]
    assert len(df) == 2
    assert df.iloc[1]["nome"] == "João"


def test_csv_ponto_e_virgula_latin1(tmp_path):
    conteudo = (
        "descrição;preço;região\n"
        "coração de alface;10,5;São Paulo\n"
        "ação promocional;3,2;Paraná\n"
        "pão francês;1,0;Ceará\n"
    )
    path = tmp_path / "dados.csv"
    path.write_bytes(conteudo.encode("latin-1"))
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["descrição", "preço", "região"]
    assert len(df) == 3
    assert df.iloc[0]["descrição"] == "coração de alface"


def test_csv_tab(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("a\tb\n1\t2\n3\t4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["a", "b"]
    assert len(df) == 2


def test_xlsx_le_primeira_aba(tmp_path):
    path = tmp_path / "dados.xlsx"
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame({"produto": ["caneta", "papel"], "valor": [1.5, 2.0]}).to_excel(
            writer, sheet_name="Primeira", index=False
        )
        pd.DataFrame({"x": [9]}).to_excel(writer, sheet_name="Segunda", index=False)
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["produto", "valor"]
    assert len(df) == 2


def _write_workbook(path, sheets, hidden=()):
    """Cria um .xlsx com as abas dadas (nome → DataFrame ou None p/ aba vazia)."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, df in sheets.items():
        sheet = workbook.create_sheet(name)
        if df is not None:
            sheet.append(list(df.columns))
            for row in df.itertuples(index=False):
                sheet.append(list(row))
        if name in hidden:
            sheet.sheet_state = "hidden"
    workbook.save(path)


def test_xlsx_capa_vazia_usa_a_unica_aba_com_dados(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_workbook(
        path,
        {
            "Capa": None,
            "Dados": pd.DataFrame({"produto": ["caneta", "papel"], "valor": [1.5, 2.0]}),
        },
    )
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["produto", "valor"]
    assert df["produto"].tolist() == ["caneta", "papel"]


def test_xlsx_capa_so_com_espacos_conta_como_vazia(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_workbook(
        path,
        {"Capa": pd.DataFrame({" ": ["   "]}), "Dados": pd.DataFrame({"x": [1, 2]})},
    )
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["x"]
    assert len(df) == 2


def test_xlsx_aba_oculta_e_ignorada(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_workbook(
        path,
        {
            "Oculta": pd.DataFrame({"segredo": [1]}),
            "Visivel": pd.DataFrame({"x": [1, 2, 3]}),
        },
        hidden=("Oculta",),
    )
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["x"]
    assert len(df) == 3


def test_xlsx_varias_abas_com_dados_usa_a_primeira(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_workbook(
        path,
        {
            "Capa": None,
            "Vendas": pd.DataFrame({"a": [1]}),
            "Custos": pd.DataFrame({"b": [2]}),
        },
    )
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["a"]


def test_xlsx_sem_dados_em_nenhuma_aba(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_workbook(path, {"Capa": None, "Oculta": pd.DataFrame({"x": [1]})}, hidden=("Oculta",))
    with pytest.raises(ParseError, match="não contém dados em nenhuma aba"):
        read_dataset_file(path, "xlsx")


def test_xlsx_corrompido_gera_erro_em_portugues(tmp_path):
    path = tmp_path / "dados.xlsx"
    path.write_bytes(b"isto nao e um xlsx")
    with pytest.raises(ParseError, match="corrompido"):
        read_dataset_file(path, "xlsx")


def test_xlsx_aba_unica_identico_ao_read_excel(tmp_path):
    """Regressão: planilha simples de uma aba produz o mesmo DataFrame de antes."""
    path = tmp_path / "dados.xlsx"
    pd.DataFrame({"produto": ["caneta", "papel"], "valor": [1.5, 2.0]}).to_excel(
        path, index=False
    )
    df = read_dataset_file(path, "xlsx")
    pd.testing.assert_frame_equal(df, pd.read_excel(path, sheet_name=0))


def test_xlsx_fixture_sales_example_identico_ao_read_excel():
    path = Path(__file__).parent / "fixtures" / "Sales_Example.xlsx"
    df = read_dataset_file(path, "xlsx")
    expected = pd.read_excel(path, sheet_name=0).dropna(how="all").reset_index(drop=True)
    pd.testing.assert_frame_equal(df, expected)


def _write_rows(path, rows):
    """Cria um .xlsx de aba única com as linhas cruas dadas (None = célula vazia)."""
    import openpyxl

    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)


def test_csv_coluna_unnamed_vazia_e_removida(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("a,,b\n1,,2\n3,,4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["a", "b"]
    assert len(df) == 2


def test_xlsx_coluna_unnamed_vazia_e_removida(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_rows(path, [["a", None, "b", None], [1, None, 2, None], [3, None, 4, None]])
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["a", "b"]
    assert df["b"].tolist() == [2, 4]


def test_coluna_unnamed_com_dados_e_mantida(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("a,,b\n1,5,2\n3,,4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["a", "Unnamed: 1", "b"]
    assert df["Unnamed: 1"].tolist()[0] == "5"


def test_coluna_nomeada_vazia_e_mantida(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("a,vazia,b\n1,,2\n3,,4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["a", "vazia", "b"]


def test_csv_titulo_em_uma_linha_acima_do_cabecalho(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text(
        "Relatório de vendas;;\nproduto;valor;regiao\ncaneta;1,5;SP\npapel;2;RJ\n",
        encoding="utf-8",
    )
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["produto", "valor", "regiao"]
    assert df["produto"].tolist() == ["caneta", "papel"]


def test_xlsx_titulo_em_uma_linha_acima_do_cabecalho(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_rows(
        path,
        [
            ["Relatório de vendas", None, None, None],
            ["produto", "valor", "regiao", None],
            ["caneta", 1.5, "SP", None],
            ["papel", 2.0, "RJ", None],
        ],
    )
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["produto", "valor", "regiao"]
    assert df["valor"].tolist() == [1.5, 2.0]
    assert len(df) == 2


def test_xlsx_titulo_com_linha_em_branco_antes_do_cabecalho(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_rows(
        path,
        [["Título", None], [None, None], ["a", "b"], [1, 2]],
    )
    df = read_dataset_file(path, "xlsx")
    assert list(df.columns) == ["a", "b"]
    assert len(df) == 1


def test_cabecalho_promovido_deduplica_nomes_repetidos(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("Título,,\nx,x,y\n1,2,3\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["x", "x.1", "y"]


def test_primeira_linha_esparsa_sem_titulo_nao_muda(tmp_path):
    """Cabeçalho legítimo + primeira linha com lacunas: nada é promovido."""
    path = tmp_path / "dados.csv"
    path.write_text("a,b,c\n1,,\n2,3,4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["a", "b", "c"]
    assert len(df) == 2


def test_titulo_com_segunda_linha_esparsa_nao_muda(tmp_path):
    """Cabeçalho vazio mas segunda linha com < 80% preenchida: comportamento antigo."""
    path = tmp_path / "dados.csv"
    path.write_text("Título,,,\na,,,\n1,2,3,4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert list(df.columns) == ["Título", "Unnamed: 1", "Unnamed: 2", "Unnamed: 3"]
    assert len(df) == 2


def test_csv_simples_identico_ao_read_csv(tmp_path):
    """Regressão: CSV simples continua igual à leitura direta do pandas."""
    path = tmp_path / "dados.csv"
    path.write_text("nome,idade\nAna,34\nJoão,28\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    pd.testing.assert_frame_equal(df, pd.read_csv(path, dtype=str))


def test_json_lista_de_objetos(tmp_path):
    path = tmp_path / "dados.json"
    path.write_text(
        json.dumps([{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]), encoding="utf-8"
    )
    df = read_dataset_file(path, "json")
    assert list(df.columns) == ["a", "b"]
    assert len(df) == 2


def test_json_que_nao_e_lista_gera_erro_em_portugues(tmp_path):
    path = tmp_path / "dados.json"
    path.write_text(json.dumps({"a": 1}), encoding="utf-8")
    with pytest.raises(ParseError, match="lista de objetos"):
        read_dataset_file(path, "json")


def test_json_invalido(tmp_path):
    path = tmp_path / "dados.json"
    path.write_text("{quebrado", encoding="utf-8")
    with pytest.raises(ParseError, match="JSON válido"):
        read_dataset_file(path, "json")


def test_arquivo_inexistente(tmp_path):
    with pytest.raises(ParseError, match="não encontrado"):
        read_dataset_file(tmp_path / "nada.csv", "csv")


def test_csv_vazio(tmp_path):
    path = tmp_path / "vazio.csv"
    path.write_text("", encoding="utf-8")
    with pytest.raises(ParseError, match="vazio"):
        read_dataset_file(path, "csv")


def test_linhas_totalmente_vazias_sao_removidas(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("a,b\n1,2\n,\n3,4\n", encoding="utf-8")
    df = read_dataset_file(path, "csv")
    assert len(df) == 2


# --- opções de layout (US-024): sheets / headerRow / transpose / combine ----


def _options(**overrides):
    base = {"sheets": [], "combine": False, "headerRow": 0, "transpose": False}
    base.update(overrides)
    return base


def _write_grid_workbook(path, sheets):
    """Cria um .xlsx a partir de grades cruas (nome → lista de linhas)."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(list(row))
    workbook.save(path)


def _duas_abas(path):
    _write_workbook(
        path,
        {
            "Primeira": pd.DataFrame({"produto": ["caneta", "papel"], "valor": [1.5, 2.0]}),
            "Segunda": pd.DataFrame({"cliente": ["Ana", "Bia", "Caio"], "idade": [30, 41, 25]}),
        },
    )


def test_opcoes_seleciona_aba_pelo_nome(tmp_path):
    path = tmp_path / "dados.xlsx"
    _duas_abas(path)
    df = read_dataset_file(path, "xlsx", _options(sheets=["Segunda"]))
    assert list(df.columns) == ["cliente", "idade"]
    assert len(df) == 3


def test_opcoes_aba_inexistente(tmp_path):
    path = tmp_path / "dados.xlsx"
    _duas_abas(path)
    with pytest.raises(ParseError, match="A aba 'Terceira' não existe"):
        read_dataset_file(path, "xlsx", _options(sheets=["Terceira"]))


def test_opcoes_sem_combine_usa_a_primeira_escolhida(tmp_path):
    path = tmp_path / "dados.xlsx"
    _duas_abas(path)
    df = read_dataset_file(path, "xlsx", _options(sheets=["Segunda", "Primeira"]))
    assert list(df.columns) == ["cliente", "idade"]
    assert "aba_origem" not in df.columns


def test_opcoes_header_row_xlsx(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_grid_workbook(
        path,
        {
            "Dados": [
                ["Relatório mensal", None, None],
                [None, None, None],
                ["Vendas 2024", None, None],
                ["produto", "valor", "regiao"],
                ["caneta", 1.5, "Sul"],
                ["papel", 2.0, "Norte"],
            ]
        },
    )
    df = read_dataset_file(path, "xlsx", _options(headerRow=3))
    assert list(df.columns) == ["produto", "valor", "regiao"]
    assert df["produto"].tolist() == ["caneta", "papel"]


def test_opcoes_header_row_csv(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text(
        "Relatório mensal;;\n;;\nVendas 2024;;\nproduto;valor;regiao\ncaneta;1,5;Sul\npapel;2,0;Norte\n",
        encoding="utf-8",
    )
    df = read_dataset_file(path, "csv", _options(headerRow=3))
    assert list(df.columns) == ["produto", "valor", "regiao"]
    assert df["regiao"].tolist() == ["Sul", "Norte"]


def _grade_transposta():
    return [
        ["indicador", "jan", "fev", "mar"],
        ["receita", 10, 20, 30],
        ["custo", 5, 6, 7],
        ["clientes", 100, 110, 120],
    ]


def test_opcoes_transpose_xlsx(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_grid_workbook(path, {"Dados": _grade_transposta()})
    df = read_dataset_file(path, "xlsx", _options(transpose=True))
    assert list(df.columns) == ["indicador", "receita", "custo", "clientes"]
    assert df["indicador"].tolist() == ["jan", "fev", "mar"]
    assert df["receita"].tolist() == [10, 20, 30]
    assert df["clientes"].tolist() == [100, 110, 120]


def test_opcoes_transpose_csv(tmp_path):
    path = tmp_path / "dados.csv"
    linhas = [";".join(str(c) for c in row) for row in _grade_transposta()]
    path.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    df = read_dataset_file(path, "csv", _options(transpose=True))
    assert list(df.columns) == ["indicador", "receita", "custo", "clientes"]
    assert df["custo"].tolist() == ["5", "6", "7"]


def test_opcoes_transpose_duplicados_ganham_sufixo(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_grid_workbook(
        path,
        {
            "Dados": [
                ["", "jan", "fev"],
                ["receita", 10, 20],
                ["receita", 5, 6],
                ["receita", 1, 2],
            ]
        },
    )
    df = read_dataset_file(path, "xlsx", _options(transpose=True))
    assert list(df.columns) == ["Unnamed: 0", "receita", "receita_2", "receita_3"]
    assert df["Unnamed: 0"].tolist() == ["jan", "fev"]
    assert df["receita_3"].tolist() == [1, 2]


def test_opcoes_transpose_com_header_row(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_grid_workbook(
        path, {"Dados": [["Painel de indicadores", None, None, None], *_grade_transposta()]}
    )
    df = read_dataset_file(path, "xlsx", _options(headerRow=1, transpose=True))
    assert list(df.columns) == ["indicador", "receita", "custo", "clientes"]
    assert len(df) == 3


def _abas_mensais(path, extra=None):
    sheets = {
        "Jan": [["produto", "valor"], ["caneta", 1.5], ["papel", 2.0]],
        # Mesmas colunas com caixa, espaço e ordem diferentes
        "Fev": [["Valor", "Produto "], [3.0, "lápis"], [4.0, "cola"]],
    }
    if extra:
        sheets.update(extra)
    _write_grid_workbook(path, sheets)


def test_opcoes_combine_concatena_e_marca_aba_origem(tmp_path):
    path = tmp_path / "dados.xlsx"
    _abas_mensais(path)
    df = read_dataset_file(path, "xlsx", _options(sheets=["Jan", "Fev"], combine=True))
    assert list(df.columns) == ["produto", "valor", "aba_origem"]
    assert df["produto"].tolist() == ["caneta", "papel", "lápis", "cola"]
    assert df["valor"].tolist() == [1.5, 2.0, 3.0, 4.0]
    assert df["aba_origem"].tolist() == ["Jan", "Jan", "Fev", "Fev"]


def test_opcoes_combine_colunas_diferentes(tmp_path):
    path = tmp_path / "dados.xlsx"
    _abas_mensais(path, {"Mar": [["produto", "quantidade"], ["caneta", 3]]})
    with pytest.raises(
        ParseError,
        match="^As abas Jan e Mar têm colunas diferentes e não podem ser combinadas\\.$",
    ):
        read_dataset_file(path, "xlsx", _options(sheets=["Jan", "Fev", "Mar"], combine=True))


def test_opcoes_combine_com_uma_aba_nao_marca_origem(tmp_path):
    path = tmp_path / "dados.xlsx"
    _abas_mensais(path)
    df = read_dataset_file(path, "xlsx", _options(sheets=["Fev"], combine=True))
    assert list(df.columns) == ["Valor", "Produto "]
    assert len(df) == 2


def test_opcoes_combine_respeita_header_row(tmp_path):
    path = tmp_path / "dados.xlsx"
    _write_grid_workbook(
        path,
        {
            "Jan": [["Janeiro"], ["produto", "valor"], ["caneta", 1.5]],
            "Fev": [["Fevereiro"], ["produto", "valor"], ["cola", 4.0]],
        },
    )
    df = read_dataset_file(path, "xlsx", _options(sheets=["Jan", "Fev"], combine=True, headerRow=1))
    assert list(df.columns) == ["produto", "valor", "aba_origem"]
    assert df["produto"].tolist() == ["caneta", "cola"]


def test_opcoes_sheets_ignoradas_no_csv(tmp_path):
    path = tmp_path / "dados.csv"
    path.write_text("nome,idade\nAna,34\n", encoding="utf-8")
    df = read_dataset_file(path, "csv", _options(sheets=["csv"]))
    pd.testing.assert_frame_equal(df, read_dataset_file(path, "csv"))


def test_opcoes_ignoradas_no_json(tmp_path):
    path = tmp_path / "dados.json"
    path.write_text(json.dumps([{"a": 1, "b": "x"}]), encoding="utf-8")
    df = read_dataset_file(path, "json", _options(headerRow=3, transpose=True))
    assert list(df.columns) == ["a", "b"]


def test_opcoes_padrao_equivalem_a_sem_opcoes(tmp_path):
    path = tmp_path / "dados.xlsx"
    _duas_abas(path)
    pd.testing.assert_frame_equal(
        read_dataset_file(path, "xlsx", _options()), read_dataset_file(path, "xlsx")
    )
    pd.testing.assert_frame_equal(
        read_dataset_file(path, "xlsx", None), read_dataset_file(path, "xlsx")
    )
